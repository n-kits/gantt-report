# -*- coding: utf-8 -*-
"""
Локальный архив задач: всё, что сборщик видел, копится в SQLite на домашнем ПК,
а каждый день, выпавший из трёхдневного окна, выгружается в Excel-файл.

    archive.sqlite   — база (%LOCALAPPDATA%\\gantt-collector): задачи, анализ, топонимы
    <archive_dir>\\ГГГГ-ММ-ДД.xlsx — рабочий день (с 04:00 до 04:00), листы «Задачи» и «Топонимы»

Каждый запуск обновляет задачи окна (статус, завершение, выполнение…); раз в сутки
и по --backfill перечитываются и закрытые дни — в архиве текущие статусы из Bitrix.
Сама база в репозиторий не попадает; на сайт уходят только строки ленты по дням,
зашифрованные (collect.write_days → ветка data, days/).
"""
from __future__ import annotations

import json
import logging
import sqlite3
from datetime import date, datetime, timedelta
from pathlib import Path

log = logging.getLogger("collector")

SCHEMA = """
CREATE TABLE IF NOT EXISTS tasks (
  id TEXT PRIMARY KEY,
  day TEXT,                -- рабочий день по регистрации (сутки с 04:00), ГГГГ-ММ-ДД
  name TEXT, status TEXT, start TEXT, finish TEXT, executor TEXT, deadline TEXT,
  project TEXT, product TEXT, works TEXT,
  url TEXT, description TEXT,
  first_seen TEXT, last_seen TEXT
);
CREATE INDEX IF NOT EXISTS tasks_day ON tasks(day);
CREATE TABLE IF NOT EXISTS analysis (
  task_id TEXT PRIMARY KEY,
  themes TEXT, conflict TEXT, sentiment TEXT, model TEXT, updated TEXT
);
CREATE TABLE IF NOT EXISTS toponyms (
  task_id TEXT, idx INTEGER,
  name TEXT, kind TEXT, macro TEXT, iso TEXT, lat REAL, lon REAL,
  src TEXT, dkm INTEGER, approx INTEGER,
  PRIMARY KEY (task_id, idx)
);
CREATE TABLE IF NOT EXISTS exports (
  day TEXT PRIMARY KEY, file TEXT, tasks INTEGER, exported_at TEXT
);
"""

TASK_COLS = ["name", "status", "start", "finish", "executor", "deadline", "project", "product", "works"]


def parse_dt(s: str) -> datetime | None:
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(s, fmt)
        except (TypeError, ValueError):
            continue
    return None


def workday(start: str, day_start_hour: int) -> str | None:
    dt = parse_dt(start)
    return (dt - timedelta(hours=day_start_hour)).date().isoformat() if dt else None


def connect(path: Path) -> sqlite3.Connection:
    db = sqlite3.connect(path)
    db.executescript(SCHEMA)
    return db


def update(db: sqlite3.Connection, cfg: dict, res: dict, geo_part: dict | None, texts: dict) -> int:
    """Задачи окна → база. texts: {id: полное описание}, если скачивали в этот запуск."""
    now = datetime.now().isoformat(timespec="seconds")
    n = 0
    with db:
        for row, t in zip(res["rows"], res.get("tasks", [])):
            vals = dict(zip(TASK_COLS, row))
            day = workday(vals["start"], cfg["day_start_hour"])
            db.execute(
                """INSERT INTO tasks (id, day, name, status, start, finish, executor, deadline, project, product, works,
                                      url, description, first_seen, last_seen)
                   VALUES (:id, :day, :name, :status, :start, :finish, :executor, :deadline, :project, :product, :works,
                           :url, :description, :now, :now)
                   ON CONFLICT(id) DO UPDATE SET
                     day=excluded.day, name=excluded.name, status=excluded.status, start=excluded.start,
                     finish=excluded.finish, executor=excluded.executor, deadline=excluded.deadline,
                     project=excluded.project, product=excluded.product, works=excluded.works, url=excluded.url,
                     description=COALESCE(excluded.description, tasks.description), last_seen=excluded.last_seen""",
                dict(vals, id=t["id"], day=day, url=t.get("url", ""), description=texts.get(t["id"]) or None, now=now))
            n += 1
        for it in (geo_part or {}).get("items", []):
            db.execute(
                """INSERT INTO analysis (task_id, themes, conflict, sentiment, model, updated) VALUES (?, ?, ?, ?, ?, ?)
                   ON CONFLICT(task_id) DO UPDATE SET themes=excluded.themes, conflict=excluded.conflict,
                     sentiment=excluded.sentiment, model=excluded.model, updated=excluded.updated""",
                (it["id"], " / ".join(it.get("themes", [])), it.get("conflict", ""), it.get("sentiment", ""),
                 it.get("model", ""), now))
            db.execute("DELETE FROM toponyms WHERE task_id = ?", (it["id"],))
            db.executemany(
                """INSERT INTO toponyms (task_id, idx, name, kind, macro, iso, lat, lon, src, dkm, approx)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                [(it["id"], i, p.get("name"), p.get("kind"), p.get("macro", ""), p.get("iso"), p.get("lat"), p.get("lon"),
                  p.get("src", ""), p.get("dkm"), 1 if p.get("approx") else 0) for i, p in enumerate(it["toponyms"])])
    return n


def prune(db: sqlite3.Connection, days: list[str], seen: set[str]) -> int:
    """Дни целиком перечитаны из Bitrix: задачи, которых там больше нет (удалены, ушли из фильтра), — убрать.
    Если Bitrix вернул пусто, а в архиве за эти дни задачи есть, — ничего не трогаем (подозрительно)."""
    if not days:
        return 0
    marks = ",".join("?" * len(days))
    ids = [i for (i,) in db.execute(f"SELECT id FROM tasks WHERE day IN ({marks})", days)]
    gone = [i for i in ids if i not in seen]
    if not gone or (not seen and ids):
        if gone:
            log.warning("Архив: Bitrix вернул 0 задач за %s, в архиве %s — не удаляю", ", ".join(days), len(ids))
        return 0
    with db:
        for tbl, col in (("tasks", "id"), ("analysis", "task_id"), ("toponyms", "task_id")):
            db.executemany(f"DELETE FROM {tbl} WHERE {col} = ?", [(i,) for i in gone])
    return len(gone)


def missing(db: sqlite3.Connection, days: list[str], seen: set[str]) -> int:
    """Сколько задач архива за эти дни Bitrix не вернул."""
    if not days:
        return 0
    marks = ",".join("?" * len(days))
    return sum(1 for (i,) in db.execute(f"SELECT id FROM tasks WHERE day IN ({marks})", days) if i not in seen)


def days_checked(db: sqlite3.Connection) -> list[tuple[str, str]]:
    """[(день, когда задачи дня последний раз сверены с Bitrix)] по возрастанию дня."""
    return db.execute("SELECT day, MAX(last_seen) FROM tasks WHERE day IS NOT NULL GROUP BY day ORDER BY day").fetchall()


def day_rows(db: sqlite3.Connection, day: str) -> list[list[str]]:
    """Строки дня в порядке полей ленты (как rows в live.json)."""
    rows = db.execute(f"SELECT {', '.join(TASK_COLS)} FROM tasks WHERE day = ?", (day,)).fetchall()
    rows.sort(key=lambda r: (parse_dt(r[2]) or datetime.min, r[0] or ""))
    return [[v if v is not None else "" for v in r] for r in rows]


# ---------------------------------------------------------------------------
# Выгрузка дня в Excel
# ---------------------------------------------------------------------------

TASK_HEADERS = ["ID", "Рабочий день", "Регистрация", "Задача", "Статус", "Завершение", "Исполнитель",
                "Крайний срок", "Проект заказчика", "Вид продукции", "Выполнение задачи",
                "Тема", "Конфликт", "Тональность", "Топонимы", "Описание", "Модель", "Последнее обновление", "Ссылка"]
TOP_HEADERS = ["ID", "Задача", "Топоним", "Тип", "Макрорегион", "Страна (ISO)", "Широта", "Долгота",
               "Источник координат", "Расхождение с оценкой, км", "Приблизительно"]
SRC_LABEL = {"cache": "кэш", "nominatim": "Nominatim", "llm": "оценка модели", "basemap": "полигон страны"}


def export_day(db: sqlite3.Connection, cfg: dict, day: str, out_dir: Path, record: bool = True) -> Path | None:
    """record=False — выгрузка вручную: день не помечается выгруженным (автовыгрузка потом перезапишет файл)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    tasks = db.execute(
        """SELECT t.id, t.day, t.start, t.name, t.status, t.finish, t.executor, t.deadline, t.project, t.product, t.works,
                  a.themes, a.conflict, a.sentiment,
                  (SELECT group_concat(name, ', ') FROM (SELECT name FROM toponyms WHERE task_id = t.id ORDER BY idx)),
                  t.description, a.model, t.last_seen, t.url
           FROM tasks t LEFT JOIN analysis a ON a.task_id = t.id
           WHERE t.day = ? ORDER BY t.start""", (day,)).fetchall()
    if not tasks:
        return None
    tasks = sorted(tasks, key=lambda r: parse_dt(r[2]) or datetime.min)
    base = cfg.get("bitrix_base", "https://crm.emg24.ru")
    tops = db.execute(
        """SELECT p.task_id, t.name, p.name, p.kind, p.macro, p.iso, p.lat, p.lon, p.src, p.dkm, p.approx
           FROM toponyms p JOIN tasks t ON t.id = p.task_id WHERE t.day = ? ORDER BY t.start, p.idx""", (day,)).fetchall()

    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F2C3A")

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(list(r))
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)

    ws = wb.active
    ws.title = "Задачи"
    sheet(ws, TASK_HEADERS,
          [list(r[:18]) + [base + r[18] if r[18] and r[18].startswith("/") else r[18]] for r in tasks],
          [10, 12, 17, 50, 15, 17, 18, 17, 18, 20, 40, 22, 18, 13, 40, 60, 18, 19, 30])
    sheet(wb.create_sheet("Топонимы"), TOP_HEADERS,
          [list(r[:8]) + [SRC_LABEL.get(r[8], r[8]), r[9], "да" if r[10] else ""] for r in tops],
          [10, 50, 28, 11, 34, 12, 10, 10, 18, 14, 14])

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{day}.xlsx"
    wb.save(path)
    if not record:
        return path
    with db:
        db.execute("""INSERT INTO exports (day, file, tasks, exported_at) VALUES (?, ?, ?, ?)
                      ON CONFLICT(day) DO UPDATE SET file=excluded.file, tasks=excluded.tasks,
                        exported_at=excluded.exported_at""",
                   (day, str(path), len(tasks), datetime.now().isoformat(timespec="seconds")))
    log.info("Архив: день %s → %s (задач %s)", day, path, len(tasks))
    return path


def export_closed(db: sqlite3.Connection, cfg: dict, window_from: datetime, out_dir: Path) -> list[Path]:
    """Выгрузить дни, целиком выпавшие из окна и ещё не выгруженные."""
    first_open = (window_from - timedelta(hours=cfg["day_start_hour"])).date().isoformat()
    days = [d for (d,) in db.execute(
        "SELECT DISTINCT day FROM tasks WHERE day < ? AND day NOT IN (SELECT day FROM exports) ORDER BY day",
        (first_open,))]
    return [p for p in (export_day(db, cfg, d, out_dir) for d in days) if p]


def reexport(db: sqlite3.Connection, cfg: dict, days: list[str], window_from: datetime, out_dir: Path) -> list[Path]:
    """Перечитанные из Bitrix закрытые дни — выгрузить заново (в Excel — текущие статусы)."""
    first_open = (window_from - timedelta(hours=cfg["day_start_hour"])).date().isoformat()
    return [p for p in (export_day(db, cfg, d, out_dir) for d in sorted(set(days)) if d < first_open) if p]


def archive_dir(cfg: dict) -> Path:
    d = cfg.get("archive_dir")
    return Path(d) if d else Path.home() / "Documents" / "Лента времени — архив"
